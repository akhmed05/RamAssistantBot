"""Офлайн-тесты: сеть заблокирована, WB подменён httpx.MockTransport."""
import asyncio
import io
import json
import socket
import sys
import unittest
from datetime import date
from pathlib import Path
from types import SimpleNamespace

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import httpx  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

from ram.bot import OwnerOnly  # noqa: E402
from ram.config import Cabinet, Settings, load_settings  # noqa: E402
from ram.db import Storage  # noqa: E402
from ram.reports import CostFileError, cost_template_xlsx, parse_cost_file  # noqa: E402
from ram.service import Manager  # noqa: E402
from ram.unit_economics import calculate  # noqa: E402
from ram.wb_client import WBClient, WBError  # noqa: E402


def _no_net(*a, **k):
    raise RuntimeError("Сеть в тестах запрещена")


socket.socket.connect = _no_net  # type: ignore[assignment]
socket.getaddrinfo = _no_net  # type: ignore[assignment]


async def _nosleep(_):
    return None


def sale(nm, qty=1, retail=1000.0, pay=750.0, sa="PLED-1"):
    return {"nm_id": nm, "sa_name": sa, "subject_name": "Пледы", "doc_type_name": "Продажа",
            "supplier_oper_name": "Продажа", "quantity": qty, "retail_amount": retail, "ppvz_for_pay": pay}


def ret(nm, qty=1, retail=1000.0, pay=750.0):
    return {"nm_id": nm, "doc_type_name": "Возврат", "supplier_oper_name": "Возврат",
            "quantity": qty, "retail_amount": retail, "ppvz_for_pay": pay}


def logistics(nm, rub):
    return {"nm_id": nm, "doc_type_name": "", "supplier_oper_name": "Логистика", "delivery_rub": rub}


REPORT = [
    sale(111, qty=1), sale(111, qty=1), sale(111, qty=1),
    ret(111),
    logistics(111, 100), logistics(111, 100), logistics(111, 100), logistics(111, 100),
    sale(222, qty=1, retail=500, pay=350, sa="FONAR"),
    logistics(222, 80),
    {"nm_id": 0, "supplier_oper_name": "Хранение", "storage_fee": 300},
    {"nm_id": 0, "supplier_oper_name": "Удержание", "deduction": 1000,
     "bonus_type_name": "Оказание услуг «ВБ.Продвижение»"},
    {"nm_id": 111, "supplier_oper_name": "Штраф", "penalty": 50},
]


class UnitEconomicsTest(unittest.TestCase):
    def test_profit_per_article(self):
        s = calculate(REPORT, costs={111: 300.0, 222: 100.0}, ad_spend={111: 200.0}, tax_percent=6)
        a = s.articles[111]
        self.assertEqual((a.sold, a.returned, a.net_qty), (3, 1, 2))
        self.assertAlmostEqual(a.revenue, 2000)
        self.assertAlmostEqual(a.to_pay, 1500)
        self.assertAlmostEqual(a.logistics, 400)
        self.assertAlmostEqual(a.penalties, 50)
        self.assertAlmostEqual(a.ads, 200)
        self.assertAlmostEqual(a.cogs, 600)
        self.assertAlmostEqual(a.tax, 120)
        # 1500 − 400 − 50 − 200 − 600 − 120
        self.assertAlmostEqual(a.profit, 130)
        self.assertAlmostEqual(a.drr, 0.1)
        b = s.articles[222]
        self.assertAlmostEqual(b.profit, 350 - 80 - 100 - 30)

    def test_ad_deduction_not_counted_twice(self):
        s = calculate(REPORT, costs={}, ad_spend={111: 200.0}, tax_percent=None)
        self.assertAlmostEqual(s.shared_costs, 300)          # только хранение
        self.assertAlmostEqual(s.shared_ad_deductions, 1000)  # не вычитается
        self.assertAlmostEqual(s.ads, 200)

    def test_without_ad_stats_uses_report_deduction(self):
        s = calculate(REPORT, costs={}, ad_spend=None, tax_percent=None)
        self.assertAlmostEqual(s.shared_costs, 1300)
        self.assertEqual(s.ads_source, "report")

    def test_missing_cost_flagged(self):
        s = calculate(REPORT, costs={111: 300.0}, ad_spend={}, tax_percent=None)
        self.assertEqual([a.nm_id for a in s.missing_cost], [222])

    def test_unmatched_ads_reduce_total(self):
        s = calculate(REPORT, costs={}, ad_spend={999: 500.0}, tax_percent=None)
        self.assertAlmostEqual(s.ads_unmatched, 500)
        total_articles = sum(a.profit for a in s.articles.values())
        self.assertAlmostEqual(s.profit, total_articles - 300 - 500)


class CostFileTest(unittest.TestCase):
    def test_template_roundtrip(self):
        data = cost_template_xlsx([(111, "PLED-1", "Пледы"), (222, "FONAR", "")], {})
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        ws["D2"] = 450
        ws["D3"] = "120,5 ₽"
        buf = io.BytesIO()
        wb.save(buf)
        items, warnings = parse_cost_file(buf.getvalue(), "x.xlsx")
        self.assertEqual(items, [(111, "PLED-1", 450.0), (222, "FONAR", 120.5)])
        self.assertEqual(warnings, [])

    def test_csv_and_bad_rows(self):
        csv = "Артикул WB;Себестоимость\n111;300\nabc;5\n222;\n".encode()
        items, warnings = parse_cost_file(csv, "c.csv")
        self.assertEqual(items, [(111, "", 300.0)])
        self.assertEqual(len(warnings), 1)

    def test_wrong_columns(self):
        wb = Workbook()
        wb.active.append(["что-то", "другое"])
        buf = io.BytesIO()
        wb.save(buf)
        with self.assertRaises(CostFileError):
            parse_cost_file(buf.getvalue(), "x.xlsx")


def wb_transport(calls):
    def handler(req: httpx.Request) -> httpx.Response:
        calls.append(req)
        assert req.method == "GET"
        assert req.headers["Authorization"] == "TOKEN-1"
        path = req.url.path
        if path.endswith("reportDetailByPeriod"):
            assert req.url.params["period"] == "daily"
            return httpx.Response(200, json=REPORT)
        if path.endswith("/promotion/count"):
            return httpx.Response(200, json={"adverts": [
                {"type": 9, "status": 9, "advert_list": [{"advertId": 5}]},
                {"type": 9, "status": 4, "advert_list": [{"advertId": 6}]},  # готова к запуску — без статистики
            ]})
        if path.endswith("/fullstats"):
            assert req.url.params["ids"] == "5"
            return httpx.Response(200, json=[{"advertId": 5, "days": [
                {"date": "2026-10-10", "apps": [{"appType": 1, "nms": [{"nmId": 111, "sum": 120.5}]}]},
                {"date": "2026-10-11", "apps": [{"appType": 1, "nms": [{"nmId": 111, "sum": 79.5}]}]},
            ]}])
        return httpx.Response(404)
    return httpx.MockTransport(handler)


class WBClientTest(unittest.TestCase):
    def test_reads_report_and_ads(self):
        calls = []

        async def go():
            c = WBClient("TOKEN-1", transport=wb_transport(calls), sleep=_nosleep)
            rows = await c.realization_report(date(2026, 10, 1), date(2026, 10, 7))
            ads = await c.ad_spend_by_article(date(2026, 10, 1), date(2026, 10, 7))
            await c.aclose()
            return rows, ads

        rows, ads = asyncio.run(go())
        self.assertEqual(len(rows), len(REPORT))
        self.assertEqual(ads, {111: 200.0})

    def test_204_is_empty_and_401_is_clear(self):
        def handler(req):
            if req.url.path.endswith("reportDetailByPeriod"):
                return httpx.Response(204)
            return httpx.Response(401, json={"detail": "secret details"})

        async def go():
            c = WBClient("T", transport=httpx.MockTransport(handler), sleep=_nosleep)
            rows = await c.realization_report(date(2026, 10, 1), date(2026, 10, 7))
            with self.assertRaises(WBError) as ctx:
                await c.campaign_ids()
            await c.aclose()
            return rows, str(ctx.exception)

        rows, msg = asyncio.run(go())
        self.assertEqual(rows, [])
        self.assertIn("токен", msg)
        self.assertNotIn("secret", msg)

    def test_429_retries(self):
        n = {"i": 0}

        def handler(req):
            n["i"] += 1
            if n["i"] == 1:
                return httpx.Response(429, headers={"X-Ratelimit-Retry": "3"})
            return httpx.Response(200, json={"adverts": []})

        async def go():
            c = WBClient("T", transport=httpx.MockTransport(handler), sleep=_nosleep)
            ids = await c.campaign_ids()
            await c.aclose()
            return ids

        self.assertEqual(asyncio.run(go()), [])
        self.assertEqual(n["i"], 2)

    def test_only_allowlisted_urls(self):
        async def go():
            c = WBClient("T", transport=httpx.MockTransport(lambda r: httpx.Response(200)), sleep=_nosleep)
            try:
                await c._get("https://example.com/x", {})
            finally:
                await c.aclose()

        with self.assertRaises(WBError):
            asyncio.run(go())


class ManagerTest(unittest.TestCase):
    def test_full_report_two_cabinets_isolated(self):
        calls = []
        settings = Settings("bot", frozenset({1}), {
            "1": Cabinet("1", "Ахмад", "TOKEN-1"),
            "2": Cabinet("2", "Марина", ""),
        })
        storage = Storage(":memory:")
        storage.save_costs("1", [(111, "PLED-1", 300.0), (222, "FONAR", 100.0)])
        storage.set_setting("1", "tax_percent", "6")
        mgr = Manager(
            settings, storage,
            client_factory=lambda tok: WBClient(tok, transport=wb_transport(calls), sleep=_nosleep),
            today=lambda: date(2026, 10, 15),
        )
        rep = asyncio.run(mgr.unit_report("1", 7))
        self.assertIn("Чистая прибыль", rep.text)
        self.assertTrue(rep.xlsx.startswith(b"PK"))
        self.assertEqual(rep.filename, "unit_1_20261008_20261014.xlsx")
        self.assertEqual(calls[0].url.params["dateTo"], "2026-10-14")
        self.assertAlmostEqual(rep.summary.articles[111].profit, 130)
        # шаблон себестоимости знает артикулы из отчёта
        self.assertEqual({a[0] for a in storage.known_articles("1")}, {111, 222})
        self.assertEqual(storage.get_costs("2"), {})
        with self.assertRaises(WBError):
            asyncio.run(mgr.unit_report("2", 7))


class AccessTest(unittest.TestCase):
    def test_owner_only(self):
        mw = OwnerOnly(frozenset({311}))
        hits = []

        async def handler(event, data):
            hits.append(event)
            return "ok"

        chat = SimpleNamespace(type="private")
        owner = SimpleNamespace(from_user=SimpleNamespace(id=311), chat=chat)
        stranger = SimpleNamespace(from_user=SimpleNamespace(id=999), chat=chat)
        group = SimpleNamespace(from_user=SimpleNamespace(id=311), chat=SimpleNamespace(type="group"))
        self.assertEqual(asyncio.run(mw(handler, owner, {})), "ok")
        self.assertIsNone(asyncio.run(mw(handler, stranger, {})))
        self.assertIsNone(asyncio.run(mw(handler, group, {})))
        self.assertEqual(len(hits), 1)

    def test_settings_parse(self):
        s = load_settings({"RAM_OWNER_IDS": "311087838, abc", "RAM_WB_2_TOKEN": " t2 ", "RAM_WB_2_NAME": "Марина"})
        self.assertEqual(s.owner_ids, frozenset({311087838}))
        self.assertFalse(s.cabinets["1"].configured)
        self.assertEqual((s.cabinets["2"].name, s.cabinets["2"].token), ("Марина", "t2"))


if __name__ == "__main__":
    unittest.main()
