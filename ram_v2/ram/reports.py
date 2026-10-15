"""Тексты для Telegram и Excel-файлы (отчёт и шаблон себестоимости)."""
from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .unit_economics import Article, Summary


def rub(x: float) -> str:
    return f"{x:,.0f} ₽".replace(",", " ")


def pct(x: float | None) -> str:
    return "—" if x is None else f"{x * 100:.1f}%"


def _name(a: Article) -> str:
    return a.vendor_code or str(a.nm_id)


def summary_text(s: Summary, cabinet_name: str, d1: date, d2: date) -> str:
    arts = list(s.articles.values())
    sold = sum(a.sold for a in arts)
    returned = sum(a.returned for a in arts)
    lines = [
        f"<b>Юнит-экономика · {cabinet_name}</b>",
        f"Период: {d1:%d.%m} – {d2:%d.%m.%Y}",
        "",
        f"Выкупы: {sold} шт, возвраты: {returned} шт",
        f"Выручка: {rub(s.revenue)}",
        f"К перечислению: {rub(sum(a.to_pay for a in arts))}",
        f"Логистика: {rub(sum(a.logistics for a in arts))}",
        f"Реклама: {rub(s.ads)} (ДРР {pct(s.ads / s.revenue if s.revenue > 0 else None)})",
        f"Себестоимость: {rub(sum(a.cogs for a in arts))}",
    ]
    storage = sum(a.storage for a in arts)
    if storage:
        lines.append(f"Хранение: {rub(storage)}")
    if s.tax_percent:
        lines.append(f"Налог {s.tax_percent:g}%: {rub(sum(a.tax for a in arts))}")
    if s.shared_costs:
        lines.append(f"Общие расходы кабинета (хранение, удержания без артикула): {rub(s.shared_costs)}")
    lines += [
        "",
        f"<b>Чистая прибыль: {rub(s.profit)}</b> (маржа {pct(s.profit / s.revenue if s.revenue > 0 else None)})",
    ]

    ranked = sorted((a for a in arts if a.revenue > 0 or a.ads > 0), key=lambda a: a.profit, reverse=True)
    if ranked:
        lines += ["", "<b>Лучшие артикулы:</b>"]
        for a in ranked[:5]:
            lines.append(f"• {_name(a)}: {rub(a.profit)}, маржа {pct(a.margin)}, ДРР {pct(a.drr)}")
        losers = [a for a in ranked if a.profit < 0][-5:]
        if losers:
            lines += ["", "<b>В минусе:</b>"]
            for a in reversed(losers):
                lines.append(f"• {_name(a)}: {rub(a.profit)}, ДРР {pct(a.drr)}")

    warnings = []
    if s.missing_cost:
        warnings.append(
            f"Нет себестоимости у {len(s.missing_cost)} арт. — прибыль по ним завышена. "
            "Загрузите себестоимость в меню «💰 Себестоимость»."
        )
    if s.tax_percent is None:
        warnings.append("Налог не учтён. Укажите ставку: /tax 6")
    if s.ads_source == "report":
        warnings.append("Статистика рекламы недоступна — реклама взята только из удержаний отчёта.")
    if s.ads_unmatched:
        warnings.append(f"Реклама по артикулам без продаж в периоде: {rub(s.ads_unmatched)} (вычтена из итога).")
    if warnings:
        lines += ["", "⚠️ " + "\n⚠️ ".join(warnings)]
    lines += ["", "Полная таблица по всем артикулам — в файле ниже."]
    return "\n".join(lines)


COLUMNS = [
    ("Артикул WB", lambda a: a.nm_id),
    ("Артикул продавца", lambda a: a.vendor_code),
    ("Предмет", lambda a: a.title),
    ("Выкуплено, шт", lambda a: a.sold),
    ("Возвраты, шт", lambda a: a.returned),
    ("Выручка", lambda a: round(a.revenue, 2)),
    ("Удержано WB с продаж", lambda a: round(a.wb_fee, 2)),
    ("К перечислению", lambda a: round(a.to_pay, 2)),
    ("Логистика", lambda a: round(a.logistics, 2)),
    ("Хранение", lambda a: round(a.storage, 2)),
    ("Штрафы", lambda a: round(a.penalties, 2)),
    ("Приёмка", lambda a: round(a.acceptance, 2)),
    ("Прочие удержания", lambda a: round(a.deductions, 2)),
    ("Доплаты", lambda a: round(a.extra_payments, 2)),
    ("Реклама", lambda a: round(a.ads, 2)),
    ("Себестоимость, шт", lambda a: a.cost_price),
    ("Себестоимость всего", lambda a: round(a.cogs, 2)),
    ("Налог", lambda a: round(a.tax, 2)),
    ("Прибыль", lambda a: round(a.profit, 2)),
    ("Прибыль на 1 шт", lambda a: None if a.profit_per_unit is None else round(a.profit_per_unit, 2)),
    ("Маржа", lambda a: a.margin),
    ("ДРР", lambda a: a.drr),
    ("ROI", lambda a: a.roi),
]
PCT_COLS = {"Маржа", "ДРР", "ROI"}


def summary_xlsx(s: Summary, cabinet_name: str, d1: date, d2: date) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Юнит-экономика"
    header_fill = PatternFill("solid", fgColor="DDE6F0")
    ws.append([c for c, _ in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for a in sorted(s.articles.values(), key=lambda a: a.profit, reverse=True):
        ws.append([fn(a) for _, fn in COLUMNS])
    for idx, (title, _) in enumerate(COLUMNS, start=1):
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = max(12, len(title) + 2)
        if title in PCT_COLS:
            for cell in ws[letter][1:]:
                cell.number_format = "0.0%"
    ws.freeze_panes = "B2"

    info = wb.create_sheet("Итог")
    rows = [
        ("Кабинет", cabinet_name),
        ("Период", f"{d1:%d.%m.%Y} – {d2:%d.%m.%Y}"),
        ("Выручка", round(s.revenue, 2)),
        ("Реклама", round(s.ads, 2)),
        ("Общие расходы без артикула", round(s.shared_costs, 2)),
        ("Реклама по артикулам без продаж", round(s.ads_unmatched, 2)),
        ("Налог, %", s.tax_percent if s.tax_percent is not None else "не указан"),
        ("Чистая прибыль", round(s.profit, 2)),
        ("", ""),
        ("Источник", "Отчёт о реализации WB (ежедневный) + статистика рекламных кампаний"),
        ("Важно", "Отчёт WB появляется с задержкой ~1 сутки; последние дни могут уточняться."),
    ]
    for r in rows:
        info.append(list(r))
    info.column_dimensions["A"].width = 34
    info.column_dimensions["B"].width = 60
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- Себестоимость ---

TEMPLATE_HEADERS = ["Артикул WB", "Артикул продавца", "Название", "Себестоимость, ₽ за 1 шт"]


def cost_template_xlsx(known: list[tuple[int, str, str]], costs: dict[int, tuple[str, float]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Себестоимость"
    ws.append(TEMPLATE_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    seen = set()
    for nm, vc, title in known:
        cost = costs.get(nm, ("", None))[1]
        ws.append([nm, vc, title, cost])
        seen.add(nm)
    for nm, (vc, cost) in costs.items():
        if nm not in seen:
            ws.append([nm, vc, "", cost])
    for letter, width in zip("ABCD", (14, 22, 30, 26)):
        ws.column_dimensions[letter].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class CostFileError(Exception):
    pass


def _num(v) -> float | None:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    text = str(v).strip().replace(" ", "").replace(" ", "").replace("₽", "").replace(",", ".")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def parse_cost_file(data: bytes, filename: str) -> tuple[list[tuple[int, str, float]], list[str]]:
    """Возвращает (строки (nm_id, артикул продавца, себестоимость), предупреждения)."""
    name = filename.lower()
    if name.endswith(".xlsx"):
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception:
            raise CostFileError("Не получилось открыть Excel-файл.") from None
        rows = [list(r) for r in wb.active.iter_rows(values_only=True)]
    elif name.endswith(".csv"):
        text = data.decode("utf-8-sig", errors="replace")
        delim = ";" if text.count(";") >= text.count(",") else ","
        rows = [line.split(delim) for line in text.splitlines()]
    else:
        raise CostFileError("Пришлите файл .xlsx (шаблон из бота) или .csv.")
    if not rows:
        raise CostFileError("Файл пустой.")

    header = [str(c or "").strip().lower() for c in rows[0]]

    def find(*needles: str) -> int | None:
        for i, h in enumerate(header):
            if any(n in h for n in needles):
                return i
        return None

    i_nm = find("артикул wb", "nm", "номенклатур")
    i_vc = find("артикул продавца", "vendor", "sa_name")
    i_cost = find("себестоим", "cost")
    if i_nm is None or i_cost is None:
        raise CostFileError(
            "Не нашёл колонки «Артикул WB» и «Себестоимость». Возьмите шаблон из бота."
        )

    items: list[tuple[int, str, float]] = []
    warnings: list[str] = []
    for n, r in enumerate(rows[1:], start=2):
        if not r or all(c in (None, "") for c in r):
            continue
        nm = _num(r[i_nm] if i_nm < len(r) else None)
        cost = _num(r[i_cost] if i_cost < len(r) else None)
        if cost is None:
            continue  # себестоимость не заполнена — пропускаем молча
        if nm is None or nm <= 0 or nm != int(nm):
            warnings.append(f"Строка {n}: неверный артикул WB")
            continue
        if cost < 0 or cost > 1_000_000:
            warnings.append(f"Строка {n}: странная себестоимость {cost}")
            continue
        vc = ""
        if i_vc is not None and i_vc < len(r) and r[i_vc] is not None:
            vc = str(r[i_vc]).strip()
        items.append((int(nm), vc, round(cost, 2)))
    if not items:
        raise CostFileError("В файле нет ни одной заполненной себестоимости.")
    return items, warnings
